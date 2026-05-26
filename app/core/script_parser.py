from __future__ import annotations

import re
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from docx import Document


SLIDE_MARKER_PATTERN = r"slide\s*[:\-–—]?\s*(\d+)\b"
OPTIONAL_TITLE_PATTERN = r"(?:\s*[:.\-)–—]\s*.*?)?"

VOICEOVER_START_RE = re.compile(
    rf"^\s*-{{3,}}\s*VOICEOVER:\s*{SLIDE_MARKER_PATTERN}{OPTIONAL_TITLE_PATTERN}\s*-{{3,}}\s*$",
    re.IGNORECASE,
)
VOICEOVER_END_RE = re.compile(r"^\s*-{3,}\s*END\s+VOICEOVER\s*-{3,}\s*$", re.IGNORECASE)
SLIDE_HEADING_RE = re.compile(
    rf"^\s*(?:#+\s*)?{SLIDE_MARKER_PATTERN}{OPTIONAL_TITLE_PATTERN}\s*$",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class ScriptSection:
    slide_number: int
    text: str
    source: str


@dataclass(frozen=True)
class ScriptParseResult:
    sections: "OrderedDict[int, ScriptSection]"
    duplicate_slide_numbers: list[int]
    debug_messages: list[str]


def _paragraphs_from_docx(docx_path: Path, debug: Callable[[str], None] | None = None) -> list[str]:
    if debug:
        debug("DOCX parser backend: python-docx")
    document = Document(str(docx_path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(paragraph.text for paragraph in cell.paragraphs)

    return paragraphs


def _flush_section(
    sections: "OrderedDict[int, ScriptSection]",
    duplicates: list[int],
    slide_number: int | None,
    lines: list[str],
    source: str,
) -> None:
    if slide_number is None:
        return
    text = "\n".join(lines).strip()
    if not text:
        return
    if slide_number in sections:
        duplicates.append(slide_number)
    sections[slide_number] = ScriptSection(slide_number=slide_number, text=text, source=source)


def parse_script(docx_path: Path, debug_log: Callable[[str], None] | None = None) -> ScriptParseResult:
    debug_messages: list[str] = []

    def debug(message: str) -> None:
        debug_messages.append(message)
        if debug_log:
            debug_log(message)

    docx_path = Path(docx_path).expanduser().resolve(strict=False)
    debug(f"Selected DOCX path: {docx_path}")
    debug(f"DOCX exists: {docx_path.exists()}")
    if docx_path.exists():
        debug(f"DOCX file size: {docx_path.stat().st_size} bytes")

    if not docx_path.exists():
        raise FileNotFoundError(f"DOCX not found: {docx_path}")
    if docx_path.suffix.lower() != ".docx":
        raise ValueError(f"Expected a .docx file, got: {docx_path}")

    try:
        paragraphs = _paragraphs_from_docx(docx_path, debug)
    except Exception as exc:
        debug(f"DOCX parser failed: {type(exc).__name__}: {exc}")
        raise
    debug(f"DOCX paragraphs read: {len(paragraphs)}")
    non_empty = [paragraph.strip() for paragraph in paragraphs if paragraph.strip()]
    for index, paragraph in enumerate(non_empty[:5], start=1):
        debug(f"DOCX first non-empty paragraph {index}: {paragraph[:180]}")
    if not paragraphs:
        raise ValueError("The DOCX file opened, but no paragraphs were read.")

    sections: "OrderedDict[int, ScriptSection]" = OrderedDict()
    duplicates: list[int] = []
    detected_headings: list[int] = []

    current_slide: int | None = None
    current_lines: list[str] = []
    current_source = "slide-heading"
    in_voiceover_block = False

    for paragraph in paragraphs:
        start_match = VOICEOVER_START_RE.match(paragraph)
        end_match = VOICEOVER_END_RE.match(paragraph)
        heading_match = SLIDE_HEADING_RE.match(paragraph)

        if start_match:
            detected_headings.append(int(start_match.group(1)))
            _flush_section(sections, duplicates, current_slide, current_lines, current_source)
            current_slide = int(start_match.group(1))
            current_lines = []
            current_source = "voiceover-block"
            in_voiceover_block = True
            continue

        if end_match and in_voiceover_block:
            _flush_section(sections, duplicates, current_slide, current_lines, current_source)
            current_slide = None
            current_lines = []
            current_source = "slide-heading"
            in_voiceover_block = False
            continue

        if not in_voiceover_block and heading_match:
            detected_headings.append(int(heading_match.group(1)))
            _flush_section(sections, duplicates, current_slide, current_lines, current_source)
            current_slide = int(heading_match.group(1))
            current_lines = []
            current_source = "slide-heading"
            continue

        if current_slide is not None:
            current_lines.append(paragraph)

    _flush_section(sections, duplicates, current_slide, current_lines, current_source)
    debug(f"Detected slide headings: {detected_headings}")
    debug(f"Parsed script sections: {sorted(sections)}")
    if not sections:
        raise ValueError(
            "No slide-based script sections were found in the DOCX. "
            "Expected headings like 'Slide 1' or '--- VOICEOVER: Slide 1 ---'."
        )
    return ScriptParseResult(sections=sections, duplicate_slide_numbers=duplicates, debug_messages=debug_messages)
